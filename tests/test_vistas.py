"""Flujo completo a través de la interfaz web (criterio de aceptación de la Fase 1)."""
import io
import uuid
from decimal import Decimal

from django.urls import reverse
from openpyxl import load_workbook

from apps.afiliados.models import Afiliado
from apps.aportaciones import services
from apps.aportaciones.models import MovimientoAfiliado

from .conftest import dias_atras


def test_flujo_completo_afiliado_movimiento_recibo(client, tesorero, caja, tipos, hoy):
    client.force_login(tesorero)
    r = client.post(reverse("afiliados:crear"), {
        "nombres": "Ficticia", "apellidos": "De Prueba", "identidad": "9999-1999-00001", "telefono": "9999-1111",
        "direccion": "N/A", "fecha_ingreso": dias_atras(30).isoformat()})
    assert r.status_code == 302, r.context["form"].errors if r.context else r
    afiliado = Afiliado.objects.get(identidad="9999199900001")

    clave = str(uuid.uuid4())
    datos = {"clave_idempotencia": clave, "afiliado": afiliado.pk, "tipo": tipos["APO"].pk, "fecha": hoy.isoformat(),
             "monto": "1250.00", "metodo_pago": "EFECTIVO", "cuenta_fondo": caja.pk, "referencia": "", "observacion": ""}
    r1 = client.post(reverse("aportaciones:crear"), datos)
    r2 = client.post(reverse("aportaciones:crear"), datos)  # doble envío del mismo formulario
    assert r1.status_code == r2.status_code == 302
    assert MovimientoAfiliado.objects.count() == 1
    m = MovimientoAfiliado.objects.get()

    r = client.get(reverse("aportaciones:recibo", args=[m.pk]))
    contenido = r.content.decode()
    assert "L 1,250.00" in contenido and "MIL DOSCIENTOS CINCUENTA LEMPIRAS CON 00/100" in contenido

    pdf = client.get(reverse("aportaciones:recibo", args=[m.pk]) + "?formato=pdf")
    assert pdf["Content-Type"] == "application/pdf" and pdf.content.startswith(b"%PDF")

    detalle = client.get(reverse("afiliados:detalle", args=[afiliado.pk]))
    assert detalle.context["saldos"]["total"] == Decimal("1250.00")


def test_formulario_muestra_errores_del_servicio(client, tesorero, afiliado, caja, tipos, hoy):
    client.force_login(tesorero)
    r = client.post(reverse("aportaciones:crear"), {
        "clave_idempotencia": uuid.uuid4(), "afiliado": afiliado.pk, "tipo": tipos["RDEP"].pk,
        "fecha": hoy.isoformat(), "monto": "10", "metodo_pago": "EFECTIVO", "cuenta_fondo": caja.pk})
    assert r.status_code == 200
    assert "Saldo insuficiente" in str(r.context["form"].errors)
    assert MovimientoAfiliado.objects.count() == 0


def test_estado_cuenta_html_pdf_xlsx(client, tesorero, afiliado, mov):
    mov(afiliado, "APO", 1000, fecha=dias_atras(20))
    mov(afiliado, "DEP", 200, fecha=dias_atras(10))
    client.force_login(tesorero)
    url = reverse("afiliados:estado_cuenta", args=[afiliado.pk])
    r = client.get(url)
    assert r.status_code == 200 and r.context["ec"].saldo_final["total"] == Decimal("1200.00")
    assert client.get(url + "?formato=pdf").content.startswith(b"%PDF")
    xlsx = client.get(url + "?formato=xlsx")
    hoja = load_workbook(io.BytesIO(xlsx.content)).active
    ultima = [c.value for c in hoja[hoja.max_row]]
    assert ultima[-1] == 1200.0


def test_reporte_movimientos_coincide_con_origen(client, tesorero, supervisor, nuevo_afiliado, mov):
    a, b = nuevo_afiliado(), nuevo_afiliado()
    mov(a, "APO", "100.25")
    mov(b, "APO", "200.50")
    mov(b, "DEP", "80")
    mov(b, "RDEP", "30")
    anulado = mov(a, "APO", "999")
    services.anular_movimiento(anulado, motivo="Prueba de reporte con anulados", usuario=supervisor)
    client.force_login(tesorero)
    r = client.get(reverse("reportes:movimientos"))
    assert r.context["neto"] == services.totales_generales()
    assert r.context["neto"] == {"aportaciones": Decimal("300.75"), "depositos": Decimal("50.00")}
    assert r.context["anulados"]["n"] == 1

    filtrado = client.get(reverse("reportes:movimientos"), {"afiliado": b.pk, "clase": "DEPOSITO"})
    assert len(filtrado.context["movimientos"]) == 2
    xlsx = client.get(reverse("reportes:movimientos"), {"formato": "xlsx"})
    hoja = load_workbook(io.BytesIO(xlsx.content)).active
    montos = [hoja.cell(row=f, column=10).value for f in range(6, hoja.max_row + 1)]
    estados = [hoja.cell(row=f, column=9).value for f in range(6, hoja.max_row + 1)]
    assert round(sum(v for v, e in zip(montos, estados) if e == "Confirmado"), 2) == 350.75
    assert client.get(reverse("reportes:movimientos"), {"formato": "pdf"}).content.startswith(b"%PDF")


def test_reporte_afiliados(client, tesorero, nuevo_afiliado, mov):
    a = nuevo_afiliado()
    mov(a, "APO", "10.10")
    mov(a, "DEP", "0.95")
    client.force_login(tesorero)
    r = client.get(reverse("reportes:afiliados"))
    assert r.context["totales"]["total"] == Decimal("11.05")
    assert client.get(reverse("reportes:afiliados"), {"formato": "pdf"}).content.startswith(b"%PDF")
    assert client.get(reverse("reportes:afiliados"), {"formato": "xlsx"}).status_code == 200


def test_listas_busqueda_filtros_orden_paginacion(client, tesorero, nuevo_afiliado, mov):
    for i in range(30):
        nuevo_afiliado(nombres=f"Persona{i:02d}")
    client.force_login(tesorero)
    r = client.get(reverse("afiliados:lista"))
    assert r.context["page_obj"].paginator.count == 30 and len(r.context["object_list"]) == 25
    assert client.get(reverse("afiliados:lista"), {"page": 2}).context["object_list"].__len__() == 5
    assert client.get(reverse("afiliados:lista"), {"q": "Persona07"}).context["page_obj"].paginator.count == 1
    desc = client.get(reverse("afiliados:lista"), {"orden": "-codigo"}).context["object_list"]
    assert desc[0].codigo == "AF-0030"
    assert client.get(reverse("afiliados:lista"), {"estado": "INACTIVO"}).context["page_obj"].paginator.count == 0


def test_dashboard_financiero(client, tesorero, afiliado, mov):
    mov(afiliado, "APO", 500)
    client.force_login(tesorero)
    r = client.get(reverse("core:dashboard"))
    assert r.status_code == 200 and r.context["totales"]["aportaciones"] == Decimal("500.00")
    assert "L 500.00" in r.content.decode()
