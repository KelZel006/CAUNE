"""Estados financieros (reportes 15 a 22): cuadran entre sí y con los servicios de cada módulo."""
import io
from datetime import date, timedelta
from decimal import Decimal as D

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.actividades import services as acs
from apps.caja import services as cs
from apps.caja.models import CategoriaCaja
from apps.core.models import Parametro
from apps.distribucion import services as ds
from apps.prestamos import services as ps
from apps.reportes import financieros as fin
from apps.reportes.views import resumen_financiero


def _fila(estado, concepto):
    return next(f for f in estado.filas if f.concepto.startswith(concepto))


def _escenario(tesorero, admin_caune, nuevo_afiliado, mov, nuevo_prestamo, caja, banco, hoy, reserva="0"):
    """Aportaciones, depósitos, préstamo con pago, actividad, saldo inicial, gastos, transferencia y una distribución
    aprobada con una entrega pagada y otra pendiente."""
    Parametro.objects.filter(clave="distribucion.porcentaje_reserva").update(valor=reserva)
    a, b = nuevo_afiliado(nombres="Ana"), nuevo_afiliado(nombres="Beto")
    mov(a, "APO", 1000, fecha=hoy - timedelta(days=60))
    mov(b, "APO", 3000, fecha=hoy - timedelta(days=60))
    mov(a, "DEP", 700, fecha=hoy - timedelta(days=5))
    mov(a, "RDEP", 200, fecha=hoy - timedelta(days=4))
    cs.registrar_movimiento(tipo="INGRESO", fecha=hoy - timedelta(days=61), concepto="Apertura", monto=D("5000"),
                            metodo_pago="EFECTIVO", cuenta_fondo=caja,
                            categoria=CategoriaCaja.objects.get(nombre="Saldo inicial"), usuario=tesorero)
    cs.registrar_movimiento(tipo="INGRESO", fecha=hoy - timedelta(days=1), concepto="Intereses del banco",
                            monto=D("400"), metodo_pago="EFECTIVO", cuenta_fondo=caja,
                            categoria=CategoriaCaja.objects.get(nombre="Intereses bancarios"), usuario=tesorero)
    cs.registrar_movimiento(tipo="EGRESO", fecha=hoy - timedelta(days=1), concepto="Papelería", monto=D("50"),
                            metodo_pago="EFECTIVO", cuenta_fondo=caja,
                            categoria=CategoriaCaja.objects.get(nombre="Gastos administrativos"), usuario=tesorero)
    p = nuevo_prestamo(b, monto="1000", plazo=2)
    ps.registrar_pago(p, usuario=tesorero, tipo="PAGO", fecha=hoy - timedelta(days=2), monto=D("507.51"),
                      metodo_pago="EFECTIVO", cuenta_fondo=caja)
    act = acs.crear_actividad({"nombre": "Rifa", "fecha_inicio": hoy - timedelta(days=5)}, tesorero)
    acs.registrar_movimiento(actividad=act, tipo="INGRESO", fecha=hoy - timedelta(days=2), concepto="Boletos",
                             monto=D("900"), metodo_pago="EFECTIVO", cuenta_fondo=caja, usuario=tesorero)
    acs.registrar_movimiento(actividad=act, tipo="GASTO", fecha=hoy - timedelta(days=2), concepto="Premio",
                             monto=D("100"), metodo_pago="EFECTIVO", cuenta_fondo=caja, usuario=tesorero)
    cs.registrar_movimiento(tipo="TRANSFERENCIA", fecha=hoy, concepto="Al banco", monto=D("2500"),
                            metodo_pago="DEPOSITO_BANCARIO", cuenta_fondo=caja, cuenta_destino=banco,
                            referencia="T-1", usuario=tesorero)
    d = ds.aprobar(ds.crear_distribucion(desde=hoy - timedelta(days=9), hasta=hoy, usuario=tesorero),
                   usuario=admin_caune)
    ds.registrar_pagos(d, afiliados=[a], forma="PAGO", fecha=hoy, cuenta_fondo=caja, metodo_pago="EFECTIVO",
                       usuario=tesorero)
    return a, b, p, d


@pytest.fixture
def escenario(sin_mora, tesorero, admin_caune, nuevo_afiliado, mov, nuevo_prestamo, caja, banco, hoy):
    return _escenario(tesorero, admin_caune, nuevo_afiliado, mov, nuevo_prestamo, caja, banco, hoy)


def test_balance_cuadra_y_coincide_con_cada_modulo(escenario, caja, banco, hoy):
    a, b, p, d = escenario
    s = fin.situacion(hoy)
    assert s["diferencia"] == D("0.00")
    assert s["efectivo"] == cs.saldo() == cs.saldo(caja) + cs.saldo(banco)
    assert s["cartera"] == s["cartera_internos"] == ps.resumen(p).saldo_capital
    assert s["depositos"] == D("500.00")
    assert s["aportaciones"] == D("4000.00")
    r = ds.resumen(d)
    pago_a = next(x["total"] for x in r["filas"] if x["afiliado"] == a)
    assert s["por_pagar"] == r["total"] - pago_a == ds.ganancias_por_pagar()
    # El total coincide con el resumen financiero (reporte 14), que no separa la cuenta por pagar.
    rf = resumen_financiero(hoy)
    assert s["activo"] == rf["recursos"] and s["pasivo_patrimonio"] == rf["obligaciones"]
    assert s["no_utilidad"] == D("5000.00") and s["reserva"] == D("0.00")
    # Toda la utilidad se distribuyó: no queda resultado del ejercicio sin repartir.
    assert s["resultado_ejercicio"] + s["resultados_anteriores"] == r["sin_asignar"]
    estado = fin.balance_general(hoy)
    assert estado.diferencia == 0 and _fila(estado, "TOTAL ACTIVO").valores == [s["activo"]]


def test_reserva_se_muestra_en_el_patrimonio(sin_mora, tesorero, admin_caune, nuevo_afiliado, mov, nuevo_prestamo,
                                             caja, banco, hoy):
    _, _, _, d = _escenario(tesorero, admin_caune, nuevo_afiliado, mov, nuevo_prestamo, caja, banco, hoy,
                            reserva="10")
    s = fin.situacion(hoy)
    assert s["reserva"] == ds.resumen(d)["reserva"] > 0
    assert s["diferencia"] == D("0.00")


def test_estado_de_resultados_es_la_utilidad_que_se_reparte(escenario, hoy):
    *_, d = escenario
    u = fin.utilidad(hoy - timedelta(days=9), hoy)
    assert u["neto"] == ds.resumen(d)["resultado"]
    # Intereses del préstamo (10.00) + rifa (900 - 100) + intereses bancarios (400) - papelería (50).
    assert u["financieros"] == ps.intereses_cobrados(hasta=hoy) + ps.mora_cobrada(hasta=hoy) == D("10.00")
    assert u["neto"] == D("10.00") + D("800.00") + D("400.00") - D("50.00")
    estado = fin.estado_resultados(hoy - timedelta(days=9), hoy)
    assert _fila(estado, "UTILIDAD (PÉRDIDA) NETA").valores == [u["neto"]]
    assert _fila(estado, "Gastos administrativos").valores == [D("50.00")]
    # El saldo inicial de caja no es ingreso.
    assert not any(f.concepto == "Saldo inicial" for f in estado.filas)


def test_resultados_comparativo_agrega_variacion(escenario, hoy):
    estado = fin.estado_resultados(hoy - timedelta(days=9), hoy, comparar=True)
    assert len(estado.columnas) == 3 and estado.columnas[-1] == "Variación"
    total = _fila(estado, "UTILIDAD (PÉRDIDA) NETA").valores
    assert total[1] == D("0.00") and total[2] == total[0]


def test_cambios_en_el_patrimonio_cuadran_con_el_balance(escenario, hoy):
    desde = hoy - timedelta(days=30)
    estado = fin.cambios_patrimonio(desde, hoy)
    assert estado.diferencia == D("0.00")
    inicial, final = estado.filas[0], estado.filas[-1]
    assert inicial.valores[-1] == fin.situacion(desde - timedelta(days=1))["patrimonio"]
    assert final.valores[-1] == fin.situacion(hoy)["patrimonio"]
    for i in range(5):  # cada columna: saldo inicial + movimientos = saldo final
        assert sum((f.valores[i] for f in estado.filas[:-1]), D(0)) == final.valores[i]


def test_flujo_clasificado_cuadra_con_el_libro_de_caja(escenario, hoy):
    desde = hoy - timedelta(days=90)
    estado = fin.flujo_clasificado(desde, hoy)
    assert estado.diferencia == D("0.00")
    assert _fila(estado, "EFECTIVO AL FINAL").valores == [cs.saldo(hasta=hoy)]
    assert _fila(estado, "Préstamos desembolsados").valores == [D("-1000.00")]
    assert _fila(estado, "Retiros de depósitos").valores == [D("-200.00")]
    assert _fila(estado, "Fondos que no son utilidad recibidos").valores == [D("5000.00")]
    assert not any("Transferencia" in f.concepto for f in estado.filas)


def test_resultados_de_ejercicios_anteriores(tesorero, admin_caune, nuevo_afiliado, mov, caja, hoy):
    anterior = date(hoy.year - 1, 6, 1)
    a = nuevo_afiliado(fecha_ingreso=date(hoy.year - 2, 1, 1))
    mov(a, "APO", 1000, fecha=anterior)
    interes = CategoriaCaja.objects.get(nombre="Intereses bancarios")
    for fecha, monto in ((anterior, "300"), (hoy, "120")):
        cs.registrar_movimiento(tipo="INGRESO", fecha=fecha, concepto="Intereses", monto=D(monto),
                                metodo_pago="EFECTIVO", cuenta_fondo=caja, categoria=interes, usuario=tesorero)
    s = fin.situacion(hoy)
    assert s["resultados_anteriores"] == D("300.00") and s["resultado_ejercicio"] == D("120.00")
    # Al distribuir el ejercicio anterior, sale de resultados anteriores y no toca el ejercicio en curso.
    ds.aprobar(ds.crear_distribucion(desde=date(hoy.year - 1, 1, 1), hasta=date(hoy.year - 1, 12, 31),
                                     usuario=tesorero), usuario=admin_caune)
    s = fin.situacion(hoy)
    assert s["resultados_anteriores"] == D("0.00") and s["resultado_ejercicio"] == D("120.00")
    assert s["por_pagar"] == D("300.00") and s["diferencia"] == D("0.00")
    # Antes de aprobarla (al 31/12 del año anterior) todavía era utilidad acumulada.
    assert fin.situacion(date(hoy.year - 1, 12, 31))["resultado_ejercicio"] == D("300.00")


def test_balance_comparativo(escenario, hoy):
    estado = fin.balance_general(hoy, comparar=hoy - timedelta(days=30))
    assert estado.columnas[-1] == "Variación"
    total = _fila(estado, "TOTAL ACTIVO").valores
    assert total[0] - total[1] == total[2]
    assert fin.situacion(hoy - timedelta(days=30))["diferencia"] == D("0.00")


def test_caja_por_cuenta_y_diaria(escenario, caja, banco, hoy):
    filas = {f["cuenta"].pk: f for f in fin.movimiento_por_cuenta(hoy - timedelta(days=3), hoy)}
    assert filas[caja.pk]["final"] == cs.saldo(caja) and filas[banco.pk]["final"] == cs.saldo(banco) == D("2500.00")
    assert filas[caja.pk]["inicial"] == cs.saldo(caja, hasta=hoy - timedelta(days=4))
    consolidado = fin.caja_diaria(None, hoy - timedelta(days=3), hoy)
    assert consolidado["final"] == cs.saldo()
    # En el consolidado la transferencia de hoy no aparece como entrada ni salida.
    hoy_fila = next(d for d in consolidado["filas"] if d["fecha"] == hoy)
    assert hoy_fila["salidas"] < D("2500")
    solo_banco = fin.caja_diaria(banco, hoy - timedelta(days=3), hoy)
    assert solo_banco["entradas"] == D("2500.00") and solo_banco["final"] == D("2500.00")


def test_antiguedad_de_la_cartera(sin_mora, nuevo_afiliado, mov, nuevo_prestamo, tesorero, caja, hoy):
    b = nuevo_afiliado()
    mov(b, "APO", 3000, fecha=hoy - timedelta(days=90))
    p = nuevo_prestamo(b, monto="1000", plazo=2)  # primera cuota vencida hace 30 días, sin pagar
    a = fin.antiguedad_cartera()
    fila = a["prestamos"][0]
    assert fila["dias"] == 30 and fila["rango"] == "1 a 30 días"
    assert a["total"] == D("1000.00") and a["indice"] == D("100.00")
    ps.registrar_pago(p, usuario=tesorero, tipo="PAGO", fecha=hoy, monto=ps.resumen(p).proxima.pendiente,
                      metodo_pago="EFECTIVO", cuenta_fondo=caja)
    a = fin.antiguedad_cartera()
    assert a["prestamos"][0]["rango"] == "Al día" and a["indice"] == D("0.00")
    estado = fin.estado_antiguedad()
    assert _fila(estado, "TOTAL CARTERA").valores[0] == ps.resumen(p).saldo_capital


@pytest.mark.parametrize("nombre", ["balance", "estado_resultados", "patrimonio", "flujo_clasificado", "indicadores",
                                    "estados_financieros", "caja_cuentas", "caja_diaria", "cartera"])
def test_reportes_en_pantalla_pdf_y_excel(client, escenario, tesorero, nombre):
    client.force_login(tesorero)
    url = reverse(f"reportes:{nombre}")
    assert client.get(url).status_code == 200
    assert client.get(url + "?formato=pdf").content.startswith(b"%PDF")
    assert client.get(url + "?formato=xlsx").status_code == 200


def test_paquete_de_estados_financieros(client, escenario, tesorero, hoy):
    Parametro.objects.filter(clave="coop.firma_presidente").update(valor="Nombre Ficticio")
    client.force_login(tesorero)
    url = reverse("reportes:estados_financieros")
    r = client.get(url)
    assert [e.titulo for e in r.context["estados"]] == ["Balance general", "Estado de resultados",
                                                        "Estado de cambios en el patrimonio",
                                                        "Estado de flujo de efectivo", "Indicadores financieros"]
    assert all(e.diferencia in (None, 0) for e in r.context["estados"])
    libro = load_workbook(io.BytesIO(client.get(url + "?formato=xlsx").content))
    assert len(libro.sheetnames) == 5
    assert client.get(url + "?formato=pdf").content.startswith(b"%PDF")
    assert client.get(reverse("reportes:balance") + "?comparar=" + (hoy - timedelta(days=30)).isoformat()) \
        .context["estados"][0].columnas[-1] == "Variación"
    assert client.get(reverse("reportes:estado_resultados") + "?comparar=on").status_code == 200
    assert "15. Balance general" in client.get(reverse("reportes:indice")).content.decode()


def test_fecha_futura_no_se_acepta(client, tesorero, hoy):
    client.force_login(tesorero)
    r = client.get(reverse("reportes:balance") + "?hasta=" + (hoy + timedelta(days=5)).isoformat())
    assert r.status_code == 200 and "hasta" in r.context["form"].errors
    assert r.context["estados"][0].subtitulo == f"Al {hoy:%d/%m/%Y}"


@pytest.mark.parametrize("usuario,vista,esperado", [
    ("presidente", "reportes:balance", 200),
    ("presidente", "reportes:estados_financieros", 200),
    ("supervisor", "reportes:estado_resultados", 200),
    ("prestamos", "reportes:balance", 403),
    ("prestamos", "reportes:cartera", 200),
    ("prestamos", "reportes:caja_diaria", 403),
    ("encargado_actividades", "reportes:patrimonio", 403),
    ("sin_rol", "reportes:flujo_clasificado", 403),
    ("presidente", "reportes:caja_cuentas", 200),
])
def test_permisos_estados_financieros(client, request, usuario, vista, esperado):
    client.force_login(request.getfixturevalue(usuario))
    assert client.get(reverse(vista)).status_code == esperado
