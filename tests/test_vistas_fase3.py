"""Flujos web de la Fase 3: caja y bancos, participación en actividades, distribución de ganancias y reportes 10-14."""
import io
import uuid
from datetime import timedelta
from decimal import Decimal as D

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.actividades import services as acs
from apps.caja import services as cs
from apps.caja.models import CategoriaCaja, MovimientoCaja
from apps.distribucion import services as ds
from apps.distribucion.models import Distribucion, PagoGanancia
from apps.prestamos import services as ps
from apps.reportes.views import resumen_financiero


def test_flujo_caja_por_la_web(client, tesorero, supervisor, afiliado, mov, caja, banco, hoy):
    mov(afiliado, "APO", 3000)
    client.force_login(tesorero)
    assert client.get(reverse("caja:indice")).status_code == 200
    form = client.get(reverse("caja:registrar") + "?tipo=EGRESO")
    categorias = list(form.context["form"].fields["categoria"].queryset)
    assert categorias and all(c.tipo == "EGRESO" for c in categorias)
    datos = {"clave_idempotencia": str(uuid.uuid4()), "categoria": categorias[0].pk, "fecha": hoy.isoformat(),
             "concepto": "Papelería", "monto": "120.50", "metodo_pago": "EFECTIVO", "cuenta_fondo": caja.pk}
    assert client.post(reverse("caja:registrar") + "?tipo=EGRESO", datos).status_code == 302
    assert client.post(reverse("caja:registrar") + "?tipo=EGRESO", datos).status_code == 302  # doble envío
    assert MovimientoCaja.objects.count() == 1
    sin_saldo = {**datos, "clave_idempotencia": str(uuid.uuid4()), "monto": "99999"}
    r = client.post(reverse("caja:registrar") + "?tipo=EGRESO", sin_saldo)
    assert r.status_code == 200 and "monto" in r.context["form"].errors
    transferencia = {"clave_idempotencia": str(uuid.uuid4()), "fecha": hoy.isoformat(), "concepto": "Depósito al banco",
                     "monto": "1000", "metodo_pago": "DEPOSITO_BANCARIO", "cuenta_fondo": caja.pk,
                     "cuenta_destino": banco.pk, "referencia": "BOL-77"}
    assert client.post(reverse("caja:registrar") + "?tipo=TRANSFERENCIA", transferencia).status_code == 302
    assert cs.saldo(caja) == D("1879.50") and cs.saldo(banco) == D("1000.00")
    m = MovimientoCaja.objects.get(tipo="EGRESO")
    assert "CIENTO VEINTE LEMPIRAS CON 50/100" in client.get(reverse("caja:detalle", args=[m.pk])).content.decode()
    client.force_login(supervisor)
    r = client.post(reverse("caja:anular", args=[m.pk]), {"motivo": "Gasto registrado dos veces"})
    assert r.status_code == 302
    assert cs.saldo(caja) == D("2000.00")


def test_categorias_por_la_web(client, tesorero):
    client.force_login(tesorero)
    r = client.post(reverse("caja:categoria_crear"), {"nombre": "Donaciones", "tipo": "INGRESO",
                                                      "afecta_utilidades": "on", "activa": "on"})
    assert r.status_code == 302 and CategoriaCaja.objects.filter(nombre="Donaciones", afecta_utilidades=True).exists()
    assert client.get(reverse("caja:categorias")).status_code == 200


def test_marcar_no_participante_por_la_web(client, tesorero, afiliado, hoy):
    act = acs.crear_actividad({"nombre": "Venta de comida", "fecha_inicio": hoy}, tesorero)
    client.force_login(tesorero)
    r = client.post(reverse("actividades:excluir", args=[act.pk]),
                    {"afiliado": afiliado.pk, "motivo": "No colaboró en esta actividad"})
    assert r.status_code == 302 and acs.excluidos_ids(act) == {afiliado.pk}
    html = client.get(reverse("actividades:detalle", args=[act.pk])).content.decode()
    assert afiliado.codigo in html and "no reciben" in html
    # Ya no se ofrece en el formulario.
    form = client.get(reverse("actividades:excluir", args=[act.pk])).context["form"]
    assert afiliado not in form.fields["afiliado"].queryset
    ex = act.exclusiones.get()
    assert client.post(reverse("actividades:incluir", args=[ex.pk]),
                       {"motivo": "Sí colaboró, se corrige"}).status_code == 302
    assert acs.excluidos_ids(act) == set()


def _socios_con_utilidad(nuevo_afiliado, mov, tesorero, caja, hoy):
    a, b = nuevo_afiliado(nombres="Ana"), nuevo_afiliado(nombres="Beto")
    mov(a, "APO", 1000, fecha=hoy - timedelta(days=60))
    mov(b, "APO", 3000, fecha=hoy - timedelta(days=60))
    cs.registrar_movimiento(tipo="INGRESO", fecha=hoy - timedelta(days=1), concepto="Intereses del banco",
                            monto=D("400"), metodo_pago="EFECTIVO", cuenta_fondo=caja,
                            categoria=CategoriaCaja.objects.get(nombre="Intereses bancarios"), usuario=tesorero)
    return a, b


def test_flujo_distribucion_por_la_web(client, tesorero, admin_caune, nuevo_afiliado, mov, caja, hoy):
    a, b = _socios_con_utilidad(nuevo_afiliado, mov, tesorero, caja, hoy)
    client.force_login(tesorero)
    r = client.post(reverse("distribucion:crear"), {"fecha_inicio": (hoy - timedelta(days=9)).isoformat(),
                                                    "fecha_fin": hoy.isoformat(), "observaciones": "Prueba"})
    assert r.status_code == 302
    d = Distribucion.objects.get()
    detalle = client.get(reverse("distribucion:detalle", args=[d.pk]))
    assert detalle.context["r"]["total"] == D("400.00")
    assert {x["afiliado"].codigo: x["total"] for x in detalle.context["r"]["filas"]} == {
        a.codigo: D("100.00"), b.codigo: D("300.00")}
    # Quien calculó no puede aprobar.
    client.post(reverse("distribucion:aprobar", args=[d.pk]))
    d.refresh_from_db()
    assert d.estado == "CALCULADA"
    client.force_login(admin_caune)
    client.post(reverse("distribucion:aprobar", args=[d.pk]))
    d.refresh_from_db()
    assert d.estado == "APROBADA"
    form = client.get(reverse("distribucion:entregar", args=[d.pk])).context["form"]
    assert set(form.initial["afiliados"]) == {a.pk, b.pk}
    r = client.post(reverse("distribucion:entregar", args=[d.pk]),
                    {"forma": "PAGO", "fecha": hoy.isoformat(), "cuenta_fondo": caja.pk, "metodo_pago": "EFECTIVO",
                     "afiliados": [a.pk]})
    assert r.status_code == 302
    r = client.post(reverse("distribucion:entregar", args=[d.pk]),
                    {"forma": "CAPITALIZACION", "fecha": hoy.isoformat(), "cuenta_fondo": caja.pk, "afiliados": [b.pk]})
    assert r.status_code == 302
    d.refresh_from_db()
    assert d.estado == "APLICADA" and PagoGanancia.objects.count() == 2
    assert client.get(reverse("distribucion:detalle", args=[d.pk]) + "?formato=pdf").content.startswith(b"%PDF")
    hoja = load_workbook(io.BytesIO(client.get(reverse("distribucion:detalle", args=[d.pk]) + "?formato=xlsx").content)).active
    valores = [[c.value for c in fila] for fila in hoja.iter_rows()]
    assert any(fila[0] == b.codigo and 300 in fila for fila in valores)
    html = client.get(reverse("afiliados:detalle", args=[b.pk])).content.decode()
    assert d.codigo in html and "Capitalizada" in html


def test_reportes_fase3_coinciden_con_servicios(client, tesorero, admin_caune, nuevo_afiliado, mov, caja, hoy):
    a, b = _socios_con_utilidad(nuevo_afiliado, mov, tesorero, caja, hoy)
    d = ds.aprobar(ds.crear_distribucion(desde=hoy - timedelta(days=9), hasta=hoy, usuario=tesorero),
                   usuario=admin_caune)
    ds.registrar_pagos(d, afiliados=[a], forma="PAGO", fecha=hoy, cuenta_fondo=caja, metodo_pago="EFECTIVO",
                       usuario=tesorero)
    client.force_login(tesorero)
    libro = client.get(reverse("reportes:libro_caja") + f"?cuenta={caja.pk}")
    assert libro.context["totales"][-1] == cs.saldo(caja) == D("4300.00")
    flujo = client.get(reverse("reportes:flujo"))
    assert flujo.context["totales"][-1] == D("4300.00")
    dist = client.get(reverse("reportes:distribuciones"))
    assert dist.context["totales"][7] == D("400.00") and dist.context["totales"][8] == D("100.00")
    resumen = client.get(reverse("reportes:resumen_financiero"))
    assert resumen.context["totales"][-1] == D("0.00")
    for nombre in ("libro_caja", "flujo", "distribuciones", "resumen_financiero"):
        assert client.get(reverse(f"reportes:{nombre}") + "?formato=pdf").content.startswith(b"%PDF")
        assert client.get(reverse(f"reportes:{nombre}") + "?formato=xlsx").status_code == 200
    r = client.get(reverse("reportes:ganancias") + f"?distribucion={d.pk}&formato=pdf")
    assert r.status_code == 302 and "formato=pdf" in r["Location"]
    assert "10. Ganancias por afiliado" in client.get(reverse("reportes:indice")).content.decode()


def test_resumen_financiero_cuadra_con_todas_las_fuentes(sin_mora, tesorero, admin_caune, supervisor, nuevo_afiliado, mov,
                                                         nuevo_prestamo, caja, banco, hoy):
    a, b = _socios_con_utilidad(nuevo_afiliado, mov, tesorero, caja, hoy)
    mov(a, "DEP", 700)
    mov(a, "RDEP", 200)
    p = nuevo_prestamo(b, monto="1000", plazo=2)
    ps.registrar_pago(p, usuario=tesorero, tipo="PAGO", fecha=hoy - timedelta(days=2), monto=D("507.51"),
                      metodo_pago="EFECTIVO", cuenta_fondo=caja)
    act = acs.crear_actividad({"nombre": "Rifa", "fecha_inicio": hoy - timedelta(days=5)}, tesorero)
    acs.registrar_movimiento(actividad=act, tipo="INGRESO", fecha=hoy - timedelta(days=2), concepto="Boletos",
                             monto=D("900"), metodo_pago="EFECTIVO", cuenta_fondo=caja, usuario=tesorero)
    acs.excluir_afiliado(act, a, usuario=tesorero, motivo="No vendió boletos")
    cs.registrar_movimiento(tipo="INGRESO", fecha=hoy - timedelta(days=30), concepto="Apertura", monto=D("5000"),
                            metodo_pago="EFECTIVO", cuenta_fondo=caja,
                            categoria=CategoriaCaja.objects.get(nombre="Saldo inicial"), usuario=tesorero)
    cs.registrar_movimiento(tipo="TRANSFERENCIA", fecha=hoy, concepto="Al banco", monto=D("2500"),
                            metodo_pago="DEPOSITO_BANCARIO", cuenta_fondo=caja, cuenta_destino=banco,
                            referencia="T-1", usuario=tesorero)
    d = ds.aprobar(ds.crear_distribucion(desde=hoy - timedelta(days=9), hasta=hoy, usuario=tesorero),
                   usuario=admin_caune)
    # A: 25 % de 410 (intereses 10 + banco 400). B: 75 % de 410 + toda la rifa (A no participa).
    totales = {x["afiliado"].codigo: x["total"] for x in ds.resumen(d)["filas"]}
    assert totales == {a.codigo: D("102.50"), b.codigo: D("1207.50")}
    ds.registrar_pagos(d, afiliados=[a], forma="PAGO", fecha=hoy, cuenta_fondo=caja, metodo_pago="EFECTIVO",
                       usuario=tesorero)
    ds.registrar_pagos(d, afiliados=[b], forma="CAPITALIZACION", fecha=hoy, cuenta_fondo=banco,
                       metodo_pago="", referencia="DIST-001", usuario=tesorero)
    r = resumen_financiero(hoy)
    assert r["diferencia"] == D("0.00")
    assert r["cartera"] == D("502.49")
    assert r["utilidades"] == D("10.00") + D("400.00") + D("900.00") - D("102.50") - D("1207.50") == D("0.00")
    assert r["no_utilidad"] == D("5000.00")


@pytest.mark.parametrize("usuario,vista,esperado", [
    ("presidente", "caja:indice", 200),
    ("presidente", "distribucion:lista", 200),
    ("presidente", "caja:registrar", 403),
    ("presidente", "distribucion:crear", 403),
    ("supervisor", "caja:indice", 200),
    ("supervisor", "caja:registrar", 403),
    ("encargado_actividades", "caja:indice", 403),
    ("encargado_actividades", "distribucion:lista", 403),
    ("prestamos", "reportes:libro_caja", 403),
    ("sin_rol", "caja:indice", 403),
    ("sin_rol", "distribucion:lista", 403),
    ("tesorero", "reportes:resumen_financiero", 200),
    ("presidente", "reportes:resumen_financiero", 200),
])
def test_permisos_fase3(client, request, usuario, vista, esperado):
    client.force_login(request.getfixturevalue(usuario))
    assert client.get(reverse(vista)).status_code == esperado


def test_presidente_no_marca_participacion(client, presidente, tesorero, afiliado, hoy):
    act = acs.crear_actividad({"nombre": "Rifa", "fecha_inicio": hoy}, tesorero)
    client.force_login(presidente)
    assert client.post(reverse("actividades:excluir", args=[act.pk]),
                       {"afiliado": afiliado.pk, "motivo": "Motivo suficiente"}).status_code == 403


def test_dashboard_muestra_dinero_en_caja(client, tesorero, afiliado, mov):
    mov(afiliado, "APO", 1500)
    client.force_login(tesorero)
    r = client.get(reverse("core:dashboard"))
    assert r.context["dinero"] == D("1500.00")
