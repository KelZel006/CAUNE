"""Flujos web de la Fase 2: préstamos, pagos, actividades y reportes 4-9."""
import io
import uuid
from datetime import timedelta
from decimal import Decimal as D

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.actividades import services as acs
from apps.prestamos import services as ps
from apps.prestamos.models import PagoPrestamo, Prestamo

# Estas pruebas verifican intereses y capital; la mora tiene sus propias pruebas en test_mora.py.
pytestmark = pytest.mark.usefixtures("sin_mora")


def test_flujo_prestamo_por_la_web(client, tesorero, afiliado, caja, hoy, fondear):
    client.force_login(tesorero)
    datos = {"tipo_prestatario": "INTERNO", "afiliado": afiliado.pk, "monto_solicitado": "1000",
             "monto_aprobado": "1000", "metodo_amortizacion": "CUOTA_FIJA", "plazo": "2", "tasa": "",
             "fecha_aprobacion": (hoy - timedelta(days=40)).isoformat(),
             "fecha_primer_vencimiento": (hoy - timedelta(days=10)).isoformat()}
    sim = client.post(reverse("prestamos:crear"), {**datos, "accion": "simular"})
    assert sim.status_code == 200 and Prestamo.objects.count() == 0
    assert [f.cuota for f in sim.context["simulacion"]] == [D("507.51"), D("507.51")]
    r = client.post(reverse("prestamos:crear"), {**datos, "accion": "guardar"})
    assert r.status_code == 302
    p = Prestamo.objects.get()
    assert p.estado == "APROBADO" and p.cuotas.count() == 2

    fondear(caja, p.monto_aprobado, hoy - timedelta(days=39))
    r = client.post(reverse("prestamos:desembolsar", args=[p.pk]),
                    {"fecha": (hoy - timedelta(days=39)).isoformat(), "metodo_pago": "EFECTIVO", "cuenta": caja.pk})
    assert r.status_code == 302
    p.refresh_from_db()
    assert p.estado == "VIGENTE"

    form = client.get(reverse("prestamos:pagar", args=[p.pk]))
    assert form.context["form"].initial["monto"] == D("507.51")
    pago = {"clave_idempotencia": str(uuid.uuid4()), "tipo": "PAGO", "fecha": hoy.isoformat(), "monto": "507.51",
            "metodo_pago": "EFECTIVO", "cuenta_fondo": caja.pk}
    assert client.post(reverse("prestamos:pagar", args=[p.pk]), pago).status_code == 302
    assert client.post(reverse("prestamos:pagar", args=[p.pk]), pago).status_code == 302  # doble envío
    assert PagoPrestamo.objects.count() == 1
    recibo = PagoPrestamo.objects.get()
    html = client.get(reverse("prestamos:pago_detalle", args=[recibo.pk])).content.decode()
    assert "L 507.51" in html and "QUINIENTOS SIETE LEMPIRAS CON 51/100" in html
    assert client.get(reverse("prestamos:pago_detalle", args=[recibo.pk]) + "?formato=pdf").content.startswith(b"%PDF")
    detalle = client.get(reverse("prestamos:detalle", args=[p.pk]))
    assert detalle.context["r"].saldo_capital == D("502.49")
    assert client.get(reverse("prestamos:detalle", args=[p.pk]) + "?formato=pdf").content.startswith(b"%PDF")
    hoja = load_workbook(io.BytesIO(client.get(reverse("prestamos:detalle", args=[p.pk]) + "?formato=xlsx").content)).active
    assert hoja.max_row == 5 + 2  # encabezado de 4 filas + títulos + 2 cuotas


def test_formulario_prestamo_externo_y_errores(client, tesorero, afiliado, hoy):
    client.force_login(tesorero)
    base = {"tipo_prestatario": "EXTERNO", "monto_solicitado": "500", "monto_aprobado": "800",
            "metodo_amortizacion": "CAPITAL_FIJO", "plazo": "2", "fecha_aprobacion": hoy.isoformat(),
            "fecha_primer_vencimiento": (hoy + timedelta(days=30)).isoformat(), "accion": "guardar"}
    r = client.post(reverse("prestamos:crear"), base)
    errores = r.context["form"].errors
    assert r.status_code == 200 and "monto_aprobado" in errores and "externo_nombre" in errores
    assert "afiliado_responsable" in errores
    ok = client.post(reverse("prestamos:crear"), {**base, "monto_aprobado": "500",
                                                  "externo_nombre": "Persona Externa Ficticia",
                                                  "externo_identidad": "8888199000002",
                                                  "afiliado_responsable": afiliado.pk})
    assert ok.status_code == 302 and Prestamo.objects.get().tasa == D("5.00")


def test_flujo_actividad_por_la_web(client, tesorero, supervisor, caja, hoy):
    client.force_login(tesorero)
    r = client.post(reverse("actividades:crear"), {"nombre": "Bingo familiar", "fecha_inicio": hoy.isoformat()})
    assert r.status_code == 302
    a = acs.Actividad.objects.get()
    client.post(reverse("actividades:subactividad_crear", args=[a.pk]), {"nombre": "Comida"})
    sub = a.subactividades.get()
    for tipo, monto, s in (("INGRESO", "2000", sub.pk), ("GASTO", "600", sub.pk), ("GASTO", "150", "")):
        r = client.post(reverse("actividades:registrar", args=[a.pk]), {
            "clave_idempotencia": str(uuid.uuid4()), "tipo": tipo, "subactividad": s, "fecha": hoy.isoformat(),
            "concepto": f"{tipo} de prueba", "monto": monto, "metodo_pago": "EFECTIVO", "cuenta_fondo": caja.pk})
        assert r.status_code == 302
    detalle = client.get(reverse("actividades:detalle", args=[a.pk]))
    assert detalle.context["resultado"] == {"ingresos": D("2000.00"), "gastos": D("750.00"), "neto": D("1250.00")}
    lista = client.get(reverse("actividades:lista"))
    assert lista.context["general"]["neto"] == D("1250.00")


def test_reportes_fase2_coinciden_con_servicios(client, tesorero, afiliado, nuevo_prestamo, caja, hoy):
    p1 = nuevo_prestamo(afiliado)
    p2 = nuevo_prestamo(afiliado, monto="3000", plazo=3, metodo="CAPITAL_FIJO")
    ps.registrar_pago(p1, usuario=tesorero, tipo="PAGO", fecha=hoy, monto=D("507.51"), metodo_pago="EFECTIVO",
                      cuenta_fondo=caja)
    ps.registrar_pago(p2, usuario=tesorero, tipo="ABONO", fecha=hoy, monto=D("1000"), metodo_pago="EFECTIVO",
                      cuenta_fondo=caja)
    a = acs.crear_actividad({"nombre": "Rifa", "fecha_inicio": hoy}, tesorero)
    acs.registrar_movimiento(actividad=a, tipo="INGRESO", fecha=hoy, concepto="Boletos", monto=D("900"),
                             metodo_pago="EFECTIVO", cuenta_fondo=caja, usuario=tesorero)
    acs.registrar_movimiento(actividad=a, tipo="GASTO", fecha=hoy, concepto="Premio", monto=D("1200"),
                             metodo_pago="EFECTIVO", cuenta_fondo=caja, usuario=tesorero)
    client.force_login(tesorero)

    r = client.get(reverse("reportes:prestamos"))
    assert r.context["resumen"][1][1] == ps.cartera_vigente() == D("502.49") + D("2000.00")
    r = client.get(reverse("reportes:saldos"))
    assert r.context["totales"][3] == ps.cartera_vigente()
    r = client.get(reverse("reportes:pagos"))
    assert r.context["resumen"] == [("Total cobrado", D("1507.51")), ("Capital recuperado", D("1497.51")),
                                    ("Intereses cobrados", D("10.00")), ("Mora cobrada", D("0.00"))]
    assert r.context["resumen"][2][1] == ps.intereses_cobrados()
    r = client.get(reverse("reportes:resultados_actividades"))
    assert r.context["totales"][5] == D("-300.00") and r.context["totales"][6] == "Pérdida"
    r = client.get(reverse("reportes:actividades"))
    assert r.context["resumen"][0][1] == acs.resultado_general()["neto"]
    for nombre in ("prestamos", "saldos", "pagos", "actividades", "resultados_actividades"):
        assert client.get(reverse(f"reportes:{nombre}"), {"formato": "pdf"}).content.startswith(b"%PDF"), nombre
        assert client.get(reverse(f"reportes:{nombre}"), {"formato": "xlsx"}).status_code == 200, nombre
    cuadro = client.get(reverse("reportes:cuadro"), {"prestamo": p1.pk, "formato": "pdf"})
    assert cuadro.status_code == 302 and cuadro.url.endswith("?formato=pdf")


@pytest.mark.parametrize("usuario, vista, esperado", [
    ("presidente", "prestamos:lista", 200), ("presidente", "prestamos:crear", 403),
    ("presidente", "actividades:lista", 200), ("presidente", "actividades:crear", 403),
    ("presidente", "reportes:prestamos", 200), ("presidente", "reportes:resultados_actividades", 200),
    ("prestamos", "prestamos:crear", 200), ("prestamos", "actividades:lista", 403),
    ("encargado_actividades", "actividades:crear", 200), ("encargado_actividades", "prestamos:lista", 403),
    ("supervisor", "prestamos:crear", 403), ("supervisor", "reportes:pagos", 200),
])
def test_permisos_fase2(client, request, usuario, vista, esperado):
    client.force_login(request.getfixturevalue(usuario))
    assert client.get(reverse(vista)).status_code == esperado


def test_presidente_no_registra_pagos(client, presidente, afiliado, nuevo_prestamo, caja, hoy):
    p = nuevo_prestamo(afiliado)
    client.force_login(presidente)
    assert client.get(reverse("prestamos:detalle", args=[p.pk])).status_code == 200
    r = client.post(reverse("prestamos:pagar", args=[p.pk]), {
        "clave_idempotencia": str(uuid.uuid4()), "tipo": "PAGO", "fecha": hoy.isoformat(), "monto": "100",
        "metodo_pago": "EFECTIVO", "cuenta_fondo": caja.pk})
    assert r.status_code == 403 and PagoPrestamo.objects.count() == 0


def test_dashboard_muestra_cartera_e_intereses(client, tesorero, afiliado, nuevo_prestamo, caja, hoy):
    p = nuevo_prestamo(afiliado)
    ps.registrar_pago(p, usuario=tesorero, tipo="PAGO", fecha=hoy, monto=D("507.51"), metodo_pago="EFECTIVO",
                      cuenta_fondo=caja)
    client.force_login(tesorero)
    ctx = client.get(reverse("core:dashboard")).context
    assert ctx["cartera"] == D("502.49") and ctx["prestamos_vigentes"] == 1
    if hoy.day >= 1:
        assert ctx["intereses_mes"] == D("10.00")
