"""Exportación de documentos a PDF y hojas de cálculo."""
import io
from decimal import Decimal

from django.conf import settings
from django.contrib.staticfiles import finders
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa

from .context_processors import datos_cooperativa

FORMATO_LEMPIRAS = '"L" #,##0.00;"L" -#,##0.00'


def _recursos_pdf(uri, rel):
    """Resuelve imágenes del PDF solo desde los archivos estáticos del proyecto (nunca URLs externas)."""
    ruta = finders.find(uri.removeprefix(settings.STATIC_URL).lstrip("/"))
    if not ruta:
        raise ValueError(f"Recurso no permitido en PDF: {uri}")
    return ruta


def respuesta_pdf(plantilla, contexto, nombre_archivo, request=None):
    contexto = {"coop": datos_cooperativa(), "generado_en": timezone.localtime(), **contexto}
    html = render_to_string(plantilla, contexto, request=request)
    salida = io.BytesIO()
    resultado = pisa.CreatePDF(html, dest=salida, encoding="utf-8", link_callback=_recursos_pdf)
    if resultado.err:
        raise RuntimeError("No fue posible generar el PDF.")
    respuesta = HttpResponse(salida.getvalue(), content_type="application/pdf")
    respuesta["Content-Disposition"] = f'inline; filename="{nombre_archivo}.pdf"'
    return respuesta


def _linea_membrete(coop):
    partes = [f"{coop['nombre']} ({coop['siglas']})"] + [v for v in (
        f"RTN {coop['rtn']}" if coop.get("rtn") else "", coop.get("direccion"),
        f"Tel. {coop['telefono']}" if coop.get("telefono") else "") if v]
    return " · ".join(partes)


def _preparar_impresion(hoja, coop, titulo, fila_encabezados, columnas):
    """Configura la hoja para imprimir: carta, ajustada al ancho, encabezados repetidos, membrete y paginación."""
    hoja.page_setup.paperSize = hoja.PAPERSIZE_LETTER
    hoja.page_setup.orientation = "landscape" if columnas > 6 else "portrait"
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.print_title_rows = f"{fila_encabezados}:{fila_encabezados}"
    hoja.print_options.horizontalCentered = True
    hoja.page_margins.top, hoja.page_margins.bottom = 0.9, 0.7
    hoja.oddHeader.left.text = f"&B{coop['nombre']}&B ({coop['siglas']})"
    hoja.oddHeader.left.size = 9
    hoja.oddHeader.left.color = "13284B"
    hoja.oddHeader.right.text = titulo
    hoja.oddHeader.right.size = 9
    hoja.oddFooter.left.text = "Generado el &D a las &T"
    hoja.oddFooter.left.size = 8
    hoja.oddFooter.right.text = "Página &P de &N"
    hoja.oddFooter.right.size = 8


def respuesta_xlsx(titulo, encabezados, filas, nombre_archivo, columnas_monto=(), subtitulo="", totales=None):
    """Genera una hoja de cálculo. `columnas_monto` son índices (base 0) con formato en lempiras."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    coop = datos_cooperativa()
    hoja.append([_linea_membrete(coop)])
    hoja.append([titulo])
    hoja.append([subtitulo or f"Generado: {timezone.localtime():%d/%m/%Y %H:%M}"])
    hoja.append([])
    hoja["A1"].font = Font(bold=True, size=13)
    hoja["A2"].font = Font(bold=True, size=11)
    hoja.append(encabezados)
    fila_enc = hoja.max_row
    for celda in hoja[fila_enc]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="13284B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    if totales:
        hoja.append([float(v) if isinstance(v, Decimal) else v for v in totales])
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)
    for indice in columnas_monto:
        letra = get_column_letter(indice + 1)
        for celda in hoja[letra][fila_enc:]:
            celda.number_format = FORMATO_LEMPIRAS
    for indice in range(1, len(encabezados) + 1):
        letra = get_column_letter(indice)
        ancho = max((len(str(c.value)) for c in hoja[letra][fila_enc - 1:] if c.value is not None), default=10)
        hoja.column_dimensions[letra].width = min(max(ancho + 2, 10), 50)
    hoja.freeze_panes = hoja.cell(row=fila_enc + 1, column=1)
    _preparar_impresion(hoja, coop, titulo, fila_enc, len(encabezados))
    salida = io.BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(
        salida.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}.xlsx"'
    return respuesta


def respuesta_xlsx_estados(estados, nombre_archivo):
    """Estados financieros en Excel: una hoja por estado, con secciones, sangrías y totales en negrita."""
    libro = Workbook()
    libro.remove(libro.active)
    coop = datos_cooperativa()
    for estado in estados:
        hoja = libro.create_sheet(estado.titulo[:31])
        hoja.append([_linea_membrete(coop)])
        hoja.append([estado.titulo])
        hoja.append([estado.subtitulo])
        hoja.append([])
        hoja["A1"].font = Font(bold=True, size=13)
        hoja["A2"].font = Font(bold=True, size=11)
        hoja.append(["Concepto"] + list(estado.columnas))
        fila_enc = hoja.max_row
        for celda in hoja[fila_enc]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="13284B")
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for f in estado.filas:
            hoja.append([f.concepto] + [float(v) if isinstance(v, Decimal) else v for v in f.valores])
            fila = hoja[hoja.max_row]
            fila[0].alignment = Alignment(indent=f.nivel * 2)
            for celda, valor in zip(fila[1:], f.valores):
                if isinstance(valor, Decimal):
                    celda.number_format = FORMATO_LEMPIRAS
                else:
                    celda.alignment = Alignment(horizontal="right")
            if f.tipo in ("seccion", "subtotal", "total"):
                for celda in fila:
                    celda.font = Font(bold=True, color="13284B" if f.tipo == "seccion" else None)
            if f.tipo == "total":
                for celda in fila:
                    celda.fill = PatternFill("solid", fgColor="E6ECF6")
        if estado.nota:
            hoja.append([])
            hoja.append([estado.nota])
            hoja.cell(row=hoja.max_row, column=1).alignment = Alignment(wrap_text=False)
        hoja.column_dimensions["A"].width = 62
        for indice in range(2, len(estado.columnas) + 2):
            hoja.column_dimensions[get_column_letter(indice)].width = 20
        hoja.freeze_panes = hoja.cell(row=fila_enc + 1, column=2)
        _preparar_impresion(hoja, coop, estado.titulo, fila_enc, len(estado.columnas) + 1)
    salida = io.BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(
        salida.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}.xlsx"'
    return respuesta
